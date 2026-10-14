"""Write the Word report, the Excel register and the run audit file.

The report follows the order in the build brief:
  1. One-page summary   2. Checklist   3. Colour-coded register
  4. Opportunities and qualifications   5. Clarification questions
  6. Contingency analysis   7. Method and limitations   8. Appendix: quotes
Every page footer carries the pack's disclaimer. Every figure comes from code.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from risk_engine.charts import cumulative_curve, tornado
from risk_engine.synthesise import SYMBOLS, fmt_money, fmt_pct, top_risks

GENERAL_METHOD = [
    "The documents were read on the user's computer and split into clauses and sections. Each piece was "
    "sent to the Claude API, which proposed risk records with a verbatim quote and a judgement of "
    "likelihood and impact, each with a reason.",
    "The software checked every quote against the source text. Findings whose quote could not be found "
    "word for word were excluded from the register. References (document, clause, page) were set by the "
    "software from where each quote was actually found.",
    "Claude proposed which findings were duplicates; the software checked each grouping and merged them, "
    "keeping the figures of the strongest finding and every source quote.",
    "All scores, ratings, expected values, ratios, contingency figures and money amounts were calculated "
    "by the software from those judgements. Claude does not produce any number in this report: its "
    "narrative refers to figures through placeholders filled in by the software, and any other figure is "
    "removed.",
    "Contingency comes from a Monte Carlo simulation with a fixed random seed, so the same register "
    "always gives the same figures.",
]


# ---------------------------------------------------------------------------
# Word helpers
# ---------------------------------------------------------------------------

def _shade(cell, hex_colour: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def _page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _table(doc, headings: list[str], rows: list[list], widths_cm: list[float] | None = None, font_pt: float = 8.5):
    from docx.shared import Cm, Pt

    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    for i, h in enumerate(headings):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(font_pt)
        _shade(cell, "D9E2F3")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if value is None else str(value))
            run.font.size = Pt(font_pt)
    if widths_cm:
        table.autofit = False
        for i, w in enumerate(widths_cm):
            table.columns[i].width = Cm(w)
        for row in table.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    return table


def _para(doc, text: str, bold: bool = False, italic: bool = False, size: float | None = None, style=None):
    from docx.shared import Pt

    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    return p


# ---------------------------------------------------------------------------
# The report
# ---------------------------------------------------------------------------

def write_report(bundle: dict, pack, path: Path, work_dir: Path) -> Path:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    calc = bundle["calculation"]["result"]
    extraction = bundle["extraction"]
    opp = bundle["opportunity"]
    syn = bundle["synthesis"]
    base = calc.get("base_value")
    cur = pack.currency
    sym = SYMBOLS.get(cur, f"{cur} " if cur else "")
    colours = pack.config["rating"].get("colours", {})
    report_cfg = pack.config.get("report", {})
    decision_cfg = pack.config.get("decision", {})
    records = calc["records"]
    by_id = {r["id"]: r for r in records}

    def money_or_pct(pct, money):
        return fmt_money(money, cur) if money is not None else fmt_pct(pct)

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    a4_short, a4_long = Cm(21.0), Cm(29.7)
    section = doc.sections[0]
    section.page_width, section.page_height = a4_short, a4_long
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    run = footer.add_run(pack.footer + "    Page ")
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    _page_number(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ----- 1. Summary --------------------------------------------------------
    title = report_cfg.get("title", "Risk review")
    doc.add_heading(f"{title}: {bundle['project_name']}", level=0)
    _para(doc, f"Draft for review by a competent person. Prepared {datetime.now():%d %B %Y}. "
               f"{pack.base_value_label}: {fmt_money(base, cur) if base is not None else 'not given'}.",
          italic=True, size=9)

    doc.add_heading("1. Summary", level=1)
    indicator = syn["indicator"]
    narrative = syn.get("narrative")
    _para(doc, decision_cfg.get("title", "Decision view"), bold=True, size=11)
    if narrative:
        _para(doc, narrative["headline"], bold=True)
        _para(doc, narrative["rationale"])
    _para(doc, f"Rule-based indicator: {indicator['label'] or 'none'} ({'; '.join(indicator['reasons'])}).",
          italic=True, size=9)
    if narrative and not narrative["agrees_with_indicator"]:
        _para(doc, f"Note: the written view ({narrative['recommendation_label']}) differs from the rule-based "
                   f"indicator. Review both before deciding.", bold=True, size=9)
    if not narrative:
        _para(doc, "No written view was produced for this run"
                   + (f" ({syn['message']})." if syn.get("message") else "."), italic=True, size=9)

    _para(doc, "Top five risks", bold=True, size=11)
    summaries = {t["record_id"]: t["summary"] for t in (narrative or {}).get("top_risks", [])}
    top = top_risks(records)
    rows = [[r["id"], r["title"] + (f"\n{summaries[r['id']]}" if r["id"] in summaries else ""),
             r["rating"].upper(), money_or_pct(r["expected_value_pct"], r.get("expected_value_money"))]
            for r in top]
    t = _table(doc, ["ID", "Risk", "Rating", "Expected value"], rows, [1.4, 11.0, 1.8, 2.8])
    for i, r in enumerate(top, start=1):
        _shade(t.rows[i].cells[2], colours.get(r["rating"], "FFFFFF"))

    _para(doc, "Contingency", bold=True, size=11)
    c = calc["contingency"]
    rows = [[name, fmt_pct(c[name]["pct"]), fmt_money(c[name]["money"], cur) if c[name]["money"] is not None else ""]
            for name in ("P50", "P80") if name in c]
    _table(doc, ["Level", f"% of {pack.base_value_label.lower()}", "Amount"], rows, [3, 4, 4])
    _para(doc, f"{len(records)} risks: {calc['ratings']['red']} red, {calc['ratings']['amber']} amber, "
               f"{calc['ratings']['green']} green. Net expected value "
               f"{money_or_pct(calc['expected_value']['net_pct'], calc['expected_value']['net_money'])}.", size=9)
    if narrative and narrative["actions"]:
        _para(doc, report_cfg.get("actions_heading", "Recommended actions"), bold=True, size=11)
        for n, a in enumerate(narrative["actions"], start=1):
            _para(doc, f"{n}.  {a}")
    doc.add_page_break()

    # ----- 2. Checklist ------------------------------------------------------
    doc.add_heading("2. Checklist results", level=1)
    rows = []
    for item in extraction["checklist"]:
        a = item["answers"][0] if item["answers"] else {}
        rows.append([item["question"], item["status"], a.get("answer", ""),
                     f"{a.get('source_doc', '')}: {a.get('source_ref', '')}" if a else ""])
    _table(doc, ["Question", "Result", "Answer", "Source"], rows, [4.5, 3, 5, 4.5])

    # ----- 3. Register (landscape) ----------------------------------------------
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = a4_long, a4_short
    doc.add_heading("3. Risk register", level=1)
    _para(doc, "Red first, then amber, then green. Impacts are percentages of the "
               f"{pack.base_value_label.lower()}. Rating: likelihood x impact score; red "
               f"{pack.config['rating']['red_min']}+, amber {pack.config['rating']['amber_min']}+; any "
               "uncapped flag is red.", size=9)
    order = {"red": 0, "amber": 1, "green": 2}
    reg = sorted(records, key=lambda r: (order.get(r["rating"], 3), -r["score"], r["id"]))
    rows = []
    for r in reg:
        srcs = "; ".join(f"{s['document']}: {s['reference']}" for s in r.get("sources", []))
        rows.append([r["id"], r["type"], r["category"], r["title"], r["rating"].upper(), r["score"],
                     f"{r['likelihood']} ({r['probability']:.0%})",
                     f"{r['impact_min']:g} / {r['impact_likely']:g} / {r['impact_max']:g}",
                     money_or_pct(r["expected_value_pct"], r.get("expected_value_money")),
                     r.get("suggested_response", ""), srcs])
    t = _table(doc, ["ID", "Type", "Cat.", "Risk", "Rating", "Score", "Likelihood", "Impact % min/likely/max",
                     "Expected value", "Response", "Source"], rows,
               [1.2, 1.8, 1.1, 5.5, 1.5, 1.1, 1.8, 2.4, 2.2, 1.7, 5.4], font_pt=7.5)
    for i, r in enumerate(reg, start=1):
        _shade(t.rows[i].cells[4], colours.get(r["rating"], "FFFFFF"))

    # ----- 4. Opportunities ----------------------------------------------------
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = a4_short, a4_long
    doc.add_heading("4. Opportunities and suggested qualifications", level=1)
    if opp.get("message") and opp.get("status") != "ok":
        _para(doc, f"The opportunity finder did not complete: {opp['message']}.", italic=True)
    _para(doc, "Asymmetries (bounded downside, larger upside)", bold=True, size=11)
    if opp["asymmetries"]:
        rows = [[a["record_id"], a["description"],
                 money_or_pct(a["downside_pct"], a["downside_money"]), money_or_pct(a["upside_pct"], a["upside_money"]),
                 "no downside" if a["no_downside"] else f"{a['ratio']:.1f} : 1"] for a in opp["asymmetries"]]
        _table(doc, ["Record", "Opportunity", "Downside", "Upside", "Upside : downside"], rows, [1.5, 8, 2.3, 2.3, 2.6])
    else:
        _para(doc, "None found.", size=9)
    opportunities = [r for r in records if r["type"] == "opportunity"]
    if opportunities:
        _para(doc, "Opportunities in the register", bold=True, size=11)
        rows = [[r["id"], r["title"], money_or_pct(-r["expected_value_pct"],
                 None if r.get("expected_value_money") is None else -r["expected_value_money"])] for r in opportunities]
        _table(doc, ["ID", "Opportunity", "Expected gain"], rows, [1.5, 12, 3.2])

    _para(doc, f"Cheap mitigations (expected loss removed at least {opp['cheap_mitigation_ratio']:g} times the cost), "
               "best value first", bold=True, size=11)
    if opp["mitigations"]:
        rows = [[m["record_id"], m["action"], money_or_pct(m["cost_pct"], m["cost_money"]),
                 money_or_pct(m["expected_loss_removed_pct"], m["expected_loss_removed_money"]),
                 f"{m['ratio']:.1f} : 1", "yes" if m["cheap"] else "no"] for m in opp["mitigations"]]
        _table(doc, ["Record", "Action", "Cost", "Expected loss removed", "Value ratio", "Cheap"], rows,
               [1.5, 7.5, 2.2, 2.6, 1.8, 1.2])
    else:
        _para(doc, "None proposed.", size=9)

    _para(doc, report_cfg.get("qualifications_heading", "Suggested qualifications"), bold=True, size=11)
    if opp["qualifications"]:
        for q in opp["qualifications"]:
            doc.add_paragraph(f"{q['qualification']} ({', '.join(q['record_ids'])})", style="List Bullet")
    else:
        _para(doc, "None proposed.", size=9)

    # ----- 5. Clarifications ------------------------------------------------------
    doc.add_heading("5. " + report_cfg.get("clarifications_heading", "Clarification questions"), level=1)
    if opp["clarifications"]:
        for n, q in enumerate(opp["clarifications"], start=1):
            _para(doc, f"Q{n}.  {q['question']} ({', '.join(q['record_ids'])})")
    else:
        _para(doc, "None proposed.", size=9)

    # ----- 6. Contingency analysis ----------------------------------------------------
    doc.add_heading("6. Contingency analysis", level=1)
    rows = [["Mean" if k == "mean" else k, fmt_pct(v["pct"]),
             fmt_money(v["money"], cur) if v["money"] is not None else ""] for k, v in c.items()]
    ev = calc["expected_value"]
    rows += [["Expected value: threats", fmt_pct(ev["threats_pct"]), fmt_money(ev["threats_money"], cur) if base is not None else ""],
             ["Expected value: opportunities", fmt_pct(ev["opportunities_pct"]), fmt_money(ev["opportunities_money"], cur) if base is not None else ""],
             ["Expected value: net", fmt_pct(ev["net_pct"]), fmt_money(ev["net_money"], cur) if base is not None else ""]]
    _table(doc, ["Measure", f"% of {pack.base_value_label.lower()}", "Amount"], rows, [6, 4, 4])
    sim = calc["simulation"]
    curve = cumulative_curve(sim, base, sym, work_dir / "chart_curve.png")
    doc.add_picture(str(curve), width=Cm(16))
    _para(doc, "Read across from a chance on the left to the net risk cost below. For example, P80 is the "
               "amount the net risk cost stays at or below in 80% of simulated outcomes.", size=8.5, italic=True)
    bars = tornado(sim["contributions"], base, sym, work_dir / "chart_tornado.png")
    if bars:
        doc.add_picture(str(bars), width=Cm(16))
        _para(doc, "Each bar is the risk's average cost in the simulated outcomes close to P80, so the bars "
                   "roughly add up to the P80 figure. Savings from opportunities are shown to the left.",
              size=8.5, italic=True)
    _para(doc, "Assumptions", bold=True, size=11)
    groups = "; ".join(f"{g}: {', '.join(ids)}" for g, ids in sim["correlation_groups"].items()) or "none"
    rows = [["Iterations", sim["iterations"]], ["Random seed (fixed)", sim["seed"]],
            ["Impact distribution", f"{sim['distribution'].upper()} on minimum, most likely and maximum"],
            ["Correlation within a group", sim["correlation_in_group"]], ["Correlated groups", groups]]
    rows += [[f"Likelihood {b['score']} ({b['label']}, {b['band']})", f"{b['midpoint']:.0%} used in the simulation"]
             for b in pack.config["likelihood_scale"]]
    rows += [[f"Impact {b['score']} ({b['label']})", b["band"]] for b in pack.config["impact_scale"]]
    _table(doc, ["Assumption", "Value"], rows, [7, 9])

    # ----- 7. Method and limitations ------------------------------------------------
    doc.add_heading("7. Method and limitations", level=1)
    _para(doc, "This report is an automated risk screening. It is a draft for checking by a competent "
               "person before any decision is made, and it is not legal advice.", bold=True)
    for text in GENERAL_METHOD + list(report_cfg.get("method_notes", [])):
        doc.add_paragraph(text, style="List Bullet")
    s = extraction["summary"]
    facts = [
        f"Pieces of text reviewed: {s['pieces_reviewed']} of {s['pieces_total']}.",
        f"Findings accepted: {s['records_accepted']}; excluded (quote not verified or incomplete): {s['records_excluded']}.",
        f"Records after merging duplicates: {len(records)}.",
        f"Checklist questions answered: {s['checklist_answered']} of {s['checklist_items']}.",
    ]
    if s["quotes_verified_share"] is not None:
        facts.append(f"Share of quotes verified: {s['quotes_verified_share']:.1%}.")
    if s["pieces_not_reviewed"]:
        facts.append("Some pieces of text were NOT reviewed (declined, cut off or not sent); see the audit file. "
                     "Those parts of the documents need checking by hand.")
    if bundle["calculation"].get("excel_edits"):
        facts.append(f"The user changed {len(bundle['calculation']['excel_edits'])} value(s) in the register "
                     "before calculation; see the audit file.")
    removed = syn.get("removed_figures") or []
    facts.append(f"Figures removed from the written view because they were not calculated by the software: {len(removed)}.")
    models = bundle["models"]
    facts.append("Models used: " + ", ".join(f"{k} {v}" for k, v in models.items() if v) + ".")
    facts.append(f"API cost of this run: ${bundle['cost_usd']:.2f}.")
    for f in facts:
        doc.add_paragraph(f, style="List Bullet")
    for doc_info in extraction.get("documents", []):
        for w in doc_info.get("warnings", []):
            doc.add_paragraph(f"{doc_info['name']}: {w}", style="List Bullet")
    for sk in extraction.get("skipped_files", []):
        doc.add_paragraph(f"Not read: {sk['name']}: {sk['reason']}.", style="List Bullet")

    # ----- 8. Appendix ----------------------------------------------------------------
    doc.add_heading("8. Appendix: source quotes", level=1)
    for r in sorted(records, key=lambda r: r["id"]):
        _para(doc, f"{r['id']}  {r['title']}", bold=True, size=9.5)
        for src in r.get("sources", []):
            p = _para(doc, f"{src['document']}: {src['reference']}", italic=True, size=8.5)
            p.paragraph_format.space_after = Pt(0)
            _para(doc, f"“{src['quote']}”", size=9)

    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# The Excel register (the Stage 4 workbook plus Stage 6 sheets)
# ---------------------------------------------------------------------------

def write_register_workbook(bundle: dict, pack, path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    from risk_engine.calc_report import NOTE, write_calculation_workbook

    write_calculation_workbook(bundle["calculation"], pack.config, path)
    wb = openpyxl.load_workbook(path)
    head = PatternFill("solid", fgColor="305496")
    cur = pack.currency

    def sheet(name, headings, rows, widths):
        ws = wb.create_sheet(name)
        ws["A1"] = NOTE
        ws["A1"].font = Font(italic=True, color="7F7F7F")
        ws.append(headings)
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if rows:
            ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(len(headings))}{len(rows) + 2}"

    opp = bundle["opportunity"]
    money = f"({cur})" if cur else ""
    sheet("Checklist", ["Question", "Result", "Answer", "Document", "Reference", "Quote"],
          [[c["question"], c["status"], (c["answers"][0]["answer"] if c["answers"] else ""),
            (c["answers"][0]["source_doc"] if c["answers"] else ""), (c["answers"][0]["source_ref"] if c["answers"] else ""),
            (c["answers"][0]["quote"] if c["answers"] else "")] for c in bundle["extraction"]["checklist"]],
          [40, 22, 40, 26, 30, 60])
    sheet("Mitigations", ["Record", "Action", "Cost %", f"Cost {money}", "Expected loss removed %",
                          f"Expected loss removed {money}", "Value ratio", "Cheap", "Reason"],
          [[m["record_id"], m["action"], round(m["cost_pct"], 4), m["cost_money"], round(m["expected_loss_removed_pct"], 4),
            m["expected_loss_removed_money"], round(m["ratio"], 2), "yes" if m["cheap"] else "no", m["reason"]]
           for m in opp["mitigations"]], [8, 50, 9, 12, 12, 14, 10, 7, 40])
    sheet("Asymmetries", ["Record", "Opportunity", "Downside %", "Upside %", "Upside : downside", "Reason"],
          [[a["record_id"], a["description"], a["downside_pct"], a["upside_pct"],
            "no downside" if a["no_downside"] else round(a["ratio"], 2), a["reason"]] for a in opp["asymmetries"]],
          [8, 50, 11, 11, 14, 40])
    sheet("Clarifications", ["Question", "Records"],
          [[q["question"], ", ".join(q["record_ids"])] for q in opp["clarifications"]], [90, 20])
    sheet("Qualifications", ["Qualification", "Records"],
          [[q["qualification"], ", ".join(q["record_ids"])] for q in opp["qualifications"]], [90, 20])
    wb.save(path)
    return path
