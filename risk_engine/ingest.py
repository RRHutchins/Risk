"""Read tender documents (PDF, Word, Excel, text) into sections.

STAGE 1 VERSION: a basic reader so the dry run can count tokens.
  - PDF:   one section per page
  - Word:  the whole document as one section (paragraphs and tables)
  - Excel: one section per worksheet
  - Text:  the whole file as one section
Stage 2 replaces this with proper splitting by clause and heading.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md"}


@dataclass
class Section:
    doc: str          # file name
    ref: str          # where in the file, e.g. "page 3" or "sheet Pricing"
    text: str


@dataclass
class Document:
    name: str
    path: Path
    kind: str
    sections: list[Section] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.sections)


@dataclass
class Chunk:
    """A piece of text small enough to send to Claude in one call."""
    doc: str
    refs: list[str]
    text: str


def read_pdf(path: Path) -> Document:
    from pypdf import PdfReader

    doc = Document(path.name, path, "PDF")
    reader = PdfReader(str(path))
    blank = 0
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            blank += 1
            continue
        doc.sections.append(Section(path.name, f"page {i}", text))
    if blank:
        if blank == len(reader.pages):
            doc.warnings.append(
                "No readable text found. This is probably a scanned PDF; it cannot be "
                "reviewed until it has been converted to text (OCR)."
            )
        else:
            doc.warnings.append(f"{blank} of {len(reader.pages)} pages have no readable text (possibly scanned).")
    return doc


def read_docx(path: Path) -> Document:
    import docx

    doc = Document(path.name, path, "Word")
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n".join(lines).strip()
    if text:
        doc.sections.append(Section(path.name, "whole document", text))
    else:
        doc.warnings.append("The Word file contains no text.")
    return doc


def read_xlsx(path: Path) -> Document:
    import openpyxl

    doc = Document(path.name, path, "Excel")
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        if rows:
            doc.sections.append(Section(path.name, f"sheet {ws.title}", "\n".join(rows)))
    wb.close()
    if not doc.sections:
        doc.warnings.append("The spreadsheet contains no data.")
    return doc


def read_text(path: Path) -> Document:
    doc = Document(path.name, path, "Text")
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if text:
        doc.sections.append(Section(path.name, "whole document", text))
    return doc


READERS = {".pdf": read_pdf, ".docx": read_docx, ".xlsx": read_xlsx, ".xlsm": read_xlsx,
           ".txt": read_text, ".md": read_text}


def read_folder(folder: Path, accepted: list[str] | None = None) -> tuple[list[Document], list[str]]:
    """Read every supported file in a folder (and its sub-folders).

    Returns the documents read, and a list of messages about files skipped.
    """
    folder = Path(folder)
    if not folder.is_dir():
        raise FileNotFoundError(f"Folder not found: {folder}")
    allowed = set(accepted or SUPPORTED) & SUPPORTED
    docs, skipped = [], []
    for path in sorted(p for p in folder.rglob("*") if p.is_file()):
        if path.name.startswith(("~$", ".")):
            continue  # temporary Office lock files and hidden files
        suffix = path.suffix.lower()
        if suffix not in allowed:
            skipped.append(f"{path.name}: file type {suffix or '(none)'} is not read")
            continue
        try:
            docs.append(READERS[suffix](path))
        except Exception as e:  # a damaged file should not stop the whole run
            skipped.append(f"{path.name}: could not be read ({e.__class__.__name__}: {e})")
    return docs, skipped


def chunk_documents(docs: list[Document], max_chars: int) -> list[Chunk]:
    """Group sections into chunks of at most max_chars characters.

    Sections from the same document are combined while they fit. A section
    longer than the limit is split at paragraph (line) boundaries.
    Chunks never mix documents, so every finding has one source file.
    """
    chunks: list[Chunk] = []
    for doc in docs:
        refs: list[str] = []
        buf: list[str] = []
        size = 0

        def flush():
            nonlocal refs, buf, size
            if buf:
                chunks.append(Chunk(doc.name, refs, "\n\n".join(buf)))
            refs, buf, size = [], [], 0

        for section in doc.sections:
            for piece in _split(section.text, max_chars):
                block = f"[{section.ref}]\n{piece}"
                if size + len(block) > max_chars:
                    flush()
                if section.ref not in refs:
                    refs.append(section.ref)
                buf.append(block)
                size += len(block) + 2
        flush()
    return chunks


def _split(text: str, max_chars: int) -> list[str]:
    limit = max(200, max_chars - 50)  # leave room for the "[page n]" label
    if len(text) <= limit:
        return [text]
    pieces, current = [], ""
    for line in text.split("\n"):
        while len(line) > limit:  # a single enormous line
            if current:
                pieces.append(current)
                current = ""
            pieces.append(line[:limit])
            line = line[limit:]
        if current and len(current) + len(line) + 1 > limit:
            pieces.append(current)
            current = line
        else:
            current = f"{current}\n{line}" if current else line
    if current:
        pieces.append(current)
    return pieces
