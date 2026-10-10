"""Write the list of document sections to an Excel file for checking by eye."""

from __future__ import annotations

from pathlib import Path

from risk_engine.ingest import Document, SkippedFile

EXCEL_CELL_LIMIT = 32000  # Excel cells hold at most 32,767 characters


def write_sections_workbook(docs: list[Document], skipped: list[SkippedFile], path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sections"
    headers = ["Document", "Type", "No.", "Reference", "Clause", "Pages / rows / lines", "Characters", "Text"]
    ws.append(headers)
    for doc in docs:
        for i, s in enumerate(doc.sections, start=1):
            text = s.text if len(s.text) <= EXCEL_CELL_LIMIT else s.text[:EXCEL_CELL_LIMIT] + " [... cut short here]"
            ws.append([doc.name, doc.kind, i, s.ref, s.clause or "", s.location_span() or "", len(s.text), text])
    widths = [32, 8, 6, 45, 12, 18, 11, 100]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[:-1]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Files and warnings")
    ws2.append(["File", "Status", "Message"])
    for doc in docs:
        ws2.append([doc.name, "read", f"{len(doc.sections)} section(s)"])
        for w in doc.warnings:
            ws2.append([doc.name, "WARNING", w])
        for n in doc.notes:
            ws2.append([doc.name, "note", n])
    for s in skipped:
        ws2.append([s.name, "NOT READ", s.reason])
    for col, width in zip("ABC", (40, 10, 110)):
        ws2.column_dimensions[col].width = width
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    wb.save(path)
    return path
