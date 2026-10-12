"""Write the Stage 4 results to Excel: scored register, contingency,
contributors (tornado), cumulative curve, merges and edits.

Every figure in this workbook was calculated by code (risk_engine.calculate).
"""

from __future__ import annotations

from pathlib import Path

NOTE = ("DRAFT: automated risk screening for review by a competent person. Not legal advice. "
        "All figures are calculated by the software from the judgements in the register.")


def _money(value, currency):
    return "" if value is None else round(value, 0)


def write_calculation_workbook(audit: dict, pack_config: dict, path: Path) -> Path:
    import openpyxl
    from openpyxl.chart import BarChart, Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Font, PatternFill

    result = audit["result"]
    base = audit["base_value"]
    cur = audit.get("currency") or ""
    money_label = f"({cur})" if cur else ""
    colours = pack_config["rating"].get("colours", {})
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="305496")

    wb = openpyxl.Workbook()

    def header(ws, values, row=2):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    def note(ws):
        ws["A1"] = NOTE
        ws["A1"].font = Font(italic=True, color="7F7F7F")

    # ----- Summary -----------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    note(ws)
    merging_note = audit["merging"].get("note", "")
    rows = [
        ("Base value " + money_label, base if base is not None else "not given (percentages only)"),
        ("Records before merging", audit["records_before_merging"]),
        ("Records after merging", audit["records_after_merging"]),
        ("Duplicate merging", merging_note or f"{len(audit['merging'].get('groups', []))} group(s) merged"),
        ("Red / amber / green", f"{result['ratings']['red']} / {result['ratings']['amber']} / {result['ratings']['green']}"),
        ("", ""),
        ("Expected value", "% of base value", f"Money {money_label}"),
        ("Threats", result["expected_value"]["threats_pct"], _money(result["expected_value"]["threats_money"], cur)),
        ("Opportunities", result["expected_value"]["opportunities_pct"], _money(result["expected_value"]["opportunities_money"], cur)),
        ("Net", result["expected_value"]["net_pct"], _money(result["expected_value"]["net_money"], cur)),
        ("", ""),
        ("Contingency (Monte Carlo)", "% of base value", f"Money {money_label}"),
    ]
    for name, v in result["contingency"].items():
        rows.append((name if name != "mean" else "Mean", v["pct"], _money(v["money"], cur)))
    sim = result["simulation"]
    rows += [
        ("", ""),
        ("Simulation", ""),
        ("Iterations", sim["iterations"]),
        ("Random seed (fixed, so results repeat)", sim["seed"]),
        ("Impact distribution", sim["distribution"].upper()),
        ("Correlation within a group", sim["correlation_in_group"]),
        ("Correlation groups used", "; ".join(f"{g}: {', '.join(ids)}" for g, ids in sim["correlation_groups"].items()) or "none"),
        ("Edits taken from records.xlsx", len(audit["excel_edits"])),
        ("Claude cost of this step (USD)", audit["cost_usd"]),
        ("", ""),
        ("Note", "Expected value uses the most likely impact (as the method requires). The simulated "
                 "mean uses the whole impact range, so it is usually higher when ranges extend further "
                 "above the most likely value than below it."),
    ]
    for r in rows:
        ws.append(list(r))
    for r in ws.iter_rows(min_row=2):
        if r[0].value in ("Expected value", "Contingency (Monte Carlo)", "Simulation"):
            for c in r:
                c.font = bold
        if isinstance(r[1].value, float):
            r[1].number_format = "0.00" if abs(r[1].value) < 1000 else "#,##0"
        if len(r) > 2 and isinstance(r[2].value, (int, float)):
            r[2].number_format = "#,##0"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22

    # ----- Register ----------------------------------------------------------
    reg = wb.create_sheet("Register")
    note(reg)
    cols = ["ID", "Type", "Category", "Title", "Rating", "Score", "Likelihood (1-5)", "Probability",
            "Impact score (1-5)", "Impact min %", "Impact likely %", "Impact max %", "Expected value %",
            f"Expected value {money_label}", "Flags", "Correlation group", "Suggested response",
            "Merged from", "Sources"]
    header(reg, cols)
    order = {"red": 0, "amber": 1, "green": 2}
    records = sorted(result["records"], key=lambda r: (order.get(r["rating"], 3), -r["score"], r["id"]))
    for r in records:
        sources = "; ".join(f"{s['document']}: {s['reference']}" for s in r.get("sources", []))
        reg.append([r["id"], r["type"], r["category"], r["title"], r["rating"].upper(), r["score"], r["likelihood"],
                    r["probability"], r["impact_score"], r["impact_min"], r["impact_likely"], r["impact_max"],
                    round(r["expected_value_pct"], 4), _money(r["expected_value_money"], cur),
                    ", ".join(r.get("flags", [])), r.get("correlation_group", ""), r.get("suggested_response", ""),
                    ", ".join(r.get("merged_from", [])), sources])
        cell = reg.cell(row=reg.max_row, column=5)
        fill = colours.get(r["rating"])
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(bold=True)
    for i, w in enumerate([7, 11, 9, 34, 8, 7, 10, 10, 10, 9, 9, 9, 11, 14, 20, 15, 12, 14, 60], start=1):
        reg.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    reg.freeze_panes = "B3"
    if records:
        reg.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(len(cols))}{reg.max_row}"

    # ----- Contributors (tornado) ---------------------------------------------
    con = wb.create_sheet("Top contributors")
    note(con)
    header(con, ["ID", "Title", "Contribution to P80 %", f"Contribution to P80 {money_label}",
                 "Average cost %", "Share of iterations it occurred"])
    top = sim["contributions"][:15]
    for c in top:
        con.append([c["id"], c["title"], round(c["p80_contribution_pct"], 4),
                    _money(None if base is None else c["p80_contribution_pct"] * base / 100, cur),
                    round(c["mean_pct"], 4), round(c["occurrence_rate"], 4)])
    for i, w in enumerate([7, 40, 14, 16, 12, 14], start=1):
        con.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    if top:
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Top contributors to P80 contingency (% of base value)"
        chart.y_axis.title = "% of base value"
        chart.add_data(Reference(con, min_col=3, min_row=2, max_row=2 + len(top)), titles_from_data=True)
        chart.set_categories(Reference(con, min_col=1, min_row=3, max_row=2 + len(top)))
        chart.x_axis.scaling.orientation = "maxMin"
        chart.legend = None
        chart.height, chart.width = 10, 18
        con.add_chart(chart, "H2")

    # ----- Cumulative curve ----------------------------------------------------
    cur_ws = wb.create_sheet("Cumulative curve")
    note(cur_ws)
    header(cur_ws, ["Percentile", "Net risk cost %", f"Net risk cost {money_label}"])
    for point in sim["curve_pct"]:
        cur_ws.append([point["percentile"], round(point["value_pct"], 4),
                       _money(None if base is None else point["value_pct"] * base / 100, cur)])
    chart = ScatterChart()
    chart.title = "Chance that net risk cost is at or below a value"
    chart.style = 13
    chart.y_axis.title = "Chance (%)"
    chart.x_axis.title = "Net risk cost, % of base value"
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 100
    series = Series(Reference(cur_ws, min_col=1, min_row=3, max_row=103),
                    Reference(cur_ws, min_col=2, min_row=3, max_row=103))
    series.marker.symbol = "none"
    series.smooth = False
    chart.series.append(series)
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend = None
    chart.height, chart.width = 10, 18
    cur_ws.add_chart(chart, "E2")
    for col, w in zip("ABC", (11, 16, 18)):
        cur_ws.column_dimensions[col].width = w

    # ----- Merging and edits ----------------------------------------------------
    mg = wb.create_sheet("Merging")
    note(mg)
    header(mg, ["Records merged", "Kept as", "Combined title", "Reason", "Status"])
    by_members = {tuple(r.get("merged_from", [])): r["id"] for r in result["records"]}
    for g in audit["merging"].get("groups", []):
        kept = by_members.get(tuple(sorted(g["record_ids"])), "")
        mg.append([", ".join(g["record_ids"]), kept, g.get("title", ""), g.get("reason", ""), "merged"])
    for g in audit["merging"].get("refused", []):
        mg.append([", ".join(g.get("record_ids", [])), "", g.get("title", ""), g.get("reason", ""),
                   f"not merged: {g['problem']}"])
    if audit["merging"].get("note"):
        mg.append(["", "", "", audit["merging"]["note"], ""])
    for col, w in zip("ABCDE", (22, 10, 36, 50, 34)):
        mg.column_dimensions[col].width = w

    ed = wb.create_sheet("Your edits")
    note(ed)
    header(ed, ["Record", "Field", "Was", "Now"])
    for e in audit["excel_edits"]:
        ed.append([e["id"], e["field"], str(e["was"]), str(e["now"])])
    for col, w in zip("ABCD", (10, 22, 20, 20)):
        ed.column_dimensions[col].width = w

    wb.save(path)
    return path
