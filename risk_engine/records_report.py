"""Write extracted records to Excel for checking by eye (Stage 3).

This is a working list, not the final register: scores, colours and money
come in Stages 4 and 5.
"""

from __future__ import annotations

from pathlib import Path

NOTE = ("DRAFT: automated risk screening for review by a competent person. Not legal advice. "
        "Likelihood and impact are Claude's judgements with reasons; scores and money are added later.")

RECORD_COLUMNS = [
    ("ID", "id", 7), ("Type", "type", 11), ("Category", "category", 9), ("Title", "title", 30),
    ("Description", "description", 50), ("Document", "source_doc", 28), ("Reference", "source_ref", 34),
    ("Quote", "quote", 60), ("Likelihood (1-5)", "likelihood", 10), ("Likelihood reason", "likelihood_reason", 40),
    ("Impact min %", "impact_min", 9), ("Impact likely %", "impact_likely", 9), ("Impact max %", "impact_max", 9),
    ("Impact reason", "impact_reason", 40), ("Existing controls", "existing_controls", 30),
    ("Suggested response", "suggested_response", 12), ("Correlation group", "correlation_group", 14),
    ("Confidence", "confidence", 10), ("Confidence reason", "confidence_reason", 30), ("Flags", "flags", 18),
    ("Quote match", "quote_match", 14),
]


def write_records_workbook(processed: dict, summary: dict, pieces: list, path: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()

    def sheet(ws, columns, rows, extra=None):
        ws.append([NOTE])
        ws["A1"].font = Font(italic=True, color="7F7F7F")
        headings = [c[0] for c in columns] + ([extra[0]] if extra else [])
        ws.append(headings)
        for row in rows:
            values = [_cell(row.get(key)) for _, key, _ in columns]
            if extra:
                values.append(_cell(row.get(extra[1])))
            ws.append(values)
        widths = [c[2] for c in columns] + ([extra[2]] if extra else [])
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in ws.iter_rows(min_row=3):
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B3"
        if rows:
            ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(len(widths))}{len(rows) + 2}"

    ws = wb.active
    ws.title = "Risk records"
    sheet(ws, RECORD_COLUMNS, processed["records"])

    excluded = [dict(r, problems_text="; ".join(r.get("problems", []))) for r in processed["excluded"]]
    sheet(wb.create_sheet("Excluded"), RECORD_COLUMNS, excluded, ("Why excluded", "problems_text", 50))

    rows = []
    for c in processed["checklist"]:
        answers = c["answers"] or [{}]
        for a in answers:
            rows.append({"question": c["question"], "status": c["status"], "answer": a.get("answer", ""),
                         "source_doc": a.get("source_doc", ""), "source_ref": a.get("source_ref", ""),
                         "quote": a.get("quote", "")})
    sheet(wb.create_sheet("Checklist"),
          [("Question", "question", 40), ("Status", "status", 22), ("Answer", "answer", 45),
           ("Document", "source_doc", 28), ("Reference", "source_ref", 34), ("Quote", "quote", 60)], rows)

    piece_rows = [{"piece": p.index + 1, "doc": p.doc, "refs": "; ".join(p.refs[:6]) + (" ..." if len(p.refs) > 6 else ""),
                   "status": p.status, "message": p.message, "records": len(p.records),
                   "cost": round(p.cost_usd, 4)} for p in pieces]
    sheet(wb.create_sheet("Pieces reviewed"),
          [("Piece", "piece", 7), ("Document", "doc", 30), ("Sections", "refs", 60), ("Status", "status", 11),
           ("Message", "message", 40), ("Records", "records", 9), ("Cost (USD)", "cost", 10)], piece_rows)

    s = wb.create_sheet("Run summary")
    labels = {
        "pieces_total": "Pieces of text", "pieces_reviewed": "Pieces reviewed",
        "pieces_not_reviewed": "Pieces NOT reviewed (see 'Pieces reviewed')",
        "records_accepted": "Records accepted", "records_excluded": "Records excluded",
        "quotes_verified_share": "Share of quotes verified", "checklist_answered": "Checklist items answered",
        "checklist_items": "Checklist items", "estimated_usd": "Estimated cost (USD)",
        "actual_usd": "Actual cost (USD)", "seconds": "Run time (seconds)", "stopped_because": "Stopped because",
    }
    for key, label in labels.items():
        value = summary.get(key)
        if key == "quotes_verified_share" and value is not None:
            value = f"{value:.1%}"
        s.append([label, _cell(value)])
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 60
    wb.move_sheet("Run summary", offset=-len(wb.sheetnames) + 1)
    wb.active = 0
    wb.save(path)
    return path


def _cell(value):
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if value is None:
        return ""
    return value
