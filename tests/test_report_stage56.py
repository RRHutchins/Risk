"""Tests for the opportunity finder, synthesis and report (Stages 5 and 6)."""

import json
from pathlib import Path

import docx
import openpyxl
import pytest

from risk_engine.calc_pipeline import run_calculation
from risk_engine.config import load_settings
from risk_engine.opportunity import assess
from risk_engine.packs import load_pack
from risk_engine.report_pipeline import run_report
from risk_engine.synthesise import REMOVED, decision_indicator
from tests.fake_claude import FakeClaude, response
from tests.test_merge_and_stage4 import answers as stage4_answers
from tests.test_merge_and_stage4 import extract_demo

PACK = load_pack("construction_tender")


@pytest.fixture
def settings(tmp_path):
    s = load_settings()
    s["budget"]["ledger_file"] = str(tmp_path / "ledger.csv")
    s["output_folder"] = str(tmp_path / "outputs")
    s["simulation"]["iterations"] = 5000
    return s


def ids_in(request, tag):
    text = request["messages"][0]["content"]
    if tag == "register":
        return [json.loads(l) for l in text.splitlines() if l.startswith("{")]
    return json.loads(text.split("<facts>")[1].split("</facts>")[0])


def all_answers(doc, request):
    content = request["messages"][0]["content"]
    if "<register>" in content:
        recs = ids_in(request, "register")
        threat = next(r for r in recs if r["type"] == "threat" and "delay" in r["title"].lower())
        ground = next(r for r in recs if "ground" in r["title"].lower())
        opp = next(r for r in recs if r["type"] == "opportunity")
        return response({
            "mitigations": [
                {"record_id": threat["id"], "action": "Qualify: cap delay damages at 10% of the contract sum",
                 "cost_pct": 0.01, "residual_likelihood": 3, "residual_impact_likely": 1.0, "reason": "cap"},
                {"record_id": ground["id"], "action": "Commission further site investigation",
                 "cost_pct": 0.4, "residual_likelihood": 3, "residual_impact_likely": 2.0, "reason": "data"},
                {"record_id": ground["id"], "action": "Makes it worse", "cost_pct": 0.1,
                 "residual_likelihood": 5, "residual_impact_likely": 9.0, "reason": "bad"},
                {"record_id": opp["id"], "action": "Not a threat", "cost_pct": 0.1,
                 "residual_likelihood": 1, "residual_impact_likely": 0.1, "reason": "bad"},
            ],
            "asymmetries": [
                {"record_id": opp["id"], "description": "Re-route the water main at lower cost",
                 "downside_pct": 0.1, "upside_pct": 0.8, "reason": "provisional sum"},
                {"record_id": opp["id"], "description": "Weak one", "downside_pct": 0.5, "upside_pct": 0.6, "reason": "x"},
            ],
            "clarifications": [{"record_ids": [ground["id"]], "question": "Will the 2019 site investigation be warranted?"},
                               {"record_ids": [], "question": "Unlinked question"}],
            "qualifications": [{"record_ids": [threat["id"]], "qualification": "Our tender assumes delay damages are capped."}],
        })
    if "<facts>" in content:
        facts = ids_in(request, "facts")
        top = [t["id"] for t in facts["top_risks"]]
        return response({
            "headline": "Bid, but only with qualifications on liability and delay damages.",
            "recommendation": "conditional",
            "rationale": ("The P80 contingency is {P80_MONEY} ({P80_PCT}) with {RED_COUNT} red risks. "
                          "Uncapped liability (R004) is the main concern; margins of about 7.5% would be eroded."),
            "top_risks": [{"record_id": i, "summary": f"Summary for {i}."} for i in top],
            "actions": ["Raise the clarification on ground conditions.", "Price {EV_NET_MONEY} of expected risk."],
        })
    return stage4_answers(doc, request)


def make_report(settings, use_claude=True):
    fake = FakeClaude(all_answers)
    ex = extract_demo(settings, fake)
    calc = run_calculation(Path(ex["out_dir"]), settings, 2_500_000, client=fake, progress=lambda m: None)
    rep = run_report(Path(calc["out_dir"]), settings, client=fake, use_claude=use_claude, progress=lambda m: None)
    return rep, fake


def test_opportunity_arithmetic_is_done_by_code():
    records = [{"id": "R001", "type": "threat", "title": "t", "likelihood": 4, "impact_likely": 2.0},
               {"id": "R002", "type": "opportunity", "title": "o", "likelihood": 3, "impact_likely": 0.5}]
    out = assess({"mitigations": [{"record_id": "R001", "action": "a", "cost_pct": 0.1, "residual_likelihood": 2,
                                   "residual_impact_likely": 1.0, "reason": ""}],
                  "asymmetries": [{"record_id": "R002", "description": "d", "downside_pct": 0.2, "upside_pct": 1.0,
                                   "reason": ""}]}, records, PACK.config, 1_000_000)
    m = out["mitigations"][0]
    # before 0.45 x 2.0 = 0.90; after 0.10 x 1.0 = 0.10; removed 0.80; ratio 0.80 / 0.1 = 8
    assert m["expected_loss_removed_pct"] == pytest.approx(0.80)
    assert m["ratio"] == pytest.approx(8.0) and m["cheap"]
    assert m["expected_loss_removed_money"] == pytest.approx(8_000)
    assert out["asymmetries"][0]["ratio"] == pytest.approx(5.0)


def test_decision_indicator_rules():
    calc = {"contingency": {"P80": {"pct": 3.0}, "P50": {"pct": 1.0}}, "ratings": {"red": 0, "amber": 2, "green": 1},
            "records": [{"flags": []}]}
    assert decision_indicator(calc, [], PACK.config)["option"] == "proceed"
    calc["records"] = [{"flags": ["uncapped_liability"]}]
    ind = decision_indicator(calc, [], PACK.config)
    assert ind["option"] == "conditional" and "uncapped_liability" in ind["reasons"][0]
    calc["contingency"]["P80"]["pct"] = 16
    assert decision_indicator(calc, [], PACK.config)["option"] == "decline"


def test_full_report(settings):
    rep, fake = make_report(settings)
    b = rep["bundle"]

    # opportunity finder: invalid proposals refused with reasons, ranked by ratio
    opp = b["opportunity"]
    assert [m["action"][:7] for m in opp["mitigations"]] == ["Qualify", "Commiss"]
    assert opp["mitigations"][0]["ratio"] > opp["mitigations"][1]["ratio"]
    problems = {r["problem"] for r in opp["refused"]}
    assert {"would leave more risk than before", "mitigations apply to threats only",
            "not linked to any record in the register"} <= problems
    assert len(opp["asymmetries"]) == 1 and opp["asymmetries"][0]["ratio"] == pytest.approx(8.0)

    # synthesis: placeholders filled by code, invented figure removed, view compared with rules
    syn = b["synthesis"]
    n = syn["narrative"]
    p80 = b["calculation"]["result"]["contingency"]["P80"]["money"]
    assert f"£{p80:,.0f}" in n["rationale"]
    assert "7.5" not in n["rationale"] and REMOVED in n["rationale"]
    assert syn["removed_figures"] == [{"field": "rationale", "removed": "7.5"}]
    assert "{" not in n["actions"][1]
    assert syn["indicator"]["option"] == "conditional" and n["agrees_with_indicator"]

    # the Word report: sections in order, disclaimer in the footer of every section
    d = docx.Document(rep["report"])
    headings = [p.text for p in d.paragraphs if p.style.name.startswith("Heading 1")]
    assert headings == ["1. Summary", "2. Checklist results", "3. Risk register",
                        "4. Opportunities and suggested qualifications",
                        "5. Clarification questions for the employer", "6. Contingency analysis",
                        "7. Method and limitations", "8. Appendix: source quotes"]
    for section in d.sections:
        assert "Not legal advice." in section.footer.paragraphs[0].text
    assert len(d.inline_shapes) == 2                       # curve and tornado charts
    register_table = next(t for t in d.tables if [c.text for c in t.rows[0].cells][4:6] == ["Rating", "Score"])
    assert register_table.rows[1].cells[4].text == "RED"

    # the Excel register and the audit file
    wb = openpyxl.load_workbook(rep["register"])
    assert {"Register", "Checklist", "Mitigations", "Asymmetries", "Clarifications", "Qualifications"} <= set(wb.sheetnames)
    audit = json.loads(Path(rep["audit"]).read_text())
    synth_call = next(c for c in fake.calls if "<facts>" in c["messages"][0]["content"])
    assert synth_call["model"] == settings["models"]["synthesis"]          # Opus writes the summary
    opp_call = next(c for c in fake.calls if "<register>" in c["messages"][0]["content"])
    assert opp_call["model"] == settings["models"]["classification"]
    assert audit["models"]["synthesis"]
    assert audit["synthesis"]["prompt"].startswith("<!--")
    assert audit["cost_usd"] > 0


def test_report_without_claude(settings):
    rep, fake = make_report(settings, use_claude=False)
    b = rep["bundle"]
    assert b["synthesis"]["narrative"] is None and b["opportunity"]["mitigations"] == []
    assert b["synthesis"]["indicator"]["option"] in {"proceed", "conditional", "decline"}
    assert not any("<register>" in c["messages"][0]["content"] for c in fake.calls)
    d = docx.Document(rep["report"])
    assert any("No written view was produced" in p.text for p in d.paragraphs)
