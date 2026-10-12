"""Tests for duplicate merging and the full Stage 4 run, using a fake Claude."""

import json
from pathlib import Path

import openpyxl
import pytest

from risk_engine.calc_pipeline import CalculationError, apply_excel_edits, parse_money, run_calculation
from risk_engine.classify import apply_merges, check_groups
from risk_engine.config import PROJECT_ROOT, load_settings
from risk_engine.costs import Ledger
from risk_engine.packs import load_pack
from risk_engine.pipeline import plan_extraction, run_extraction
from tests.fake_claude import FakeClaude, record, response

DEMO = PROJECT_ROOT / "samples" / "demo_tender"
PACK = load_pack("construction_tender").config


@pytest.fixture
def settings(tmp_path):
    s = load_settings()
    s["budget"]["ledger_file"] = str(tmp_path / "ledger.csv")
    s["output_folder"] = str(tmp_path / "outputs")
    s["simulation"]["iterations"] = 5000
    return s


def r(id, type="threat", likelihood=3, likely=2.0, flags=None, doc="a.pdf"):
    return {"id": id, "type": type, "category": "GRD", "title": id, "likelihood": likelihood,
            "impact_min": likely / 2, "impact_likely": likely, "impact_max": likely * 3,
            "flags": flags or [], "source_doc": doc, "source_ref": f"ref {id}", "quote": f"quote {id}",
            "correlation_group": "ground"}


# ----- merging rules (code only) ---------------------------------------------------

def test_group_checks():
    records = [r("R001"), r("R002"), r("R003", type="opportunity"), r("R004")]
    usable, refused = check_groups([
        {"record_ids": ["R001", "R002"], "title": "Ground", "reason": "same"},
        {"record_ids": ["R002", "R004"], "title": "x", "reason": "reuses R002"},
        {"record_ids": ["R001", "R003"], "title": "x", "reason": "mixed"},
        {"record_ids": ["R004"], "title": "x", "reason": "alone"},
        {"record_ids": ["R004", "R999"], "title": "x", "reason": "unknown"},
    ], records)
    assert [g["record_ids"] for g in usable] == [["R001", "R002"]]
    assert [g["problem"] for g in refused] == ["a record is already in another group", "mixes threats and opportunities",
                                               "a group needs at least two records", "names a record that does not exist"]


def test_merge_keeps_strongest_numbers_all_sources_and_all_flags():
    records = [r("R001", likelihood=2, likely=1.0), r("R002", likelihood=4, likely=3.0, flags=["uncapped_liability"],
               doc="b.pdf"), r("R003")]
    out = apply_merges(records, [{"record_ids": ["R001", "R002"], "title": "Ground risk", "reason": "same"}], PACK)
    assert [x["id"] for x in out] == ["R002", "R003"]          # strongest member's ID is kept
    merged = out[0]
    assert merged["likelihood"] == 4 and merged["impact_likely"] == 3.0
    assert merged["title"] == "Ground risk" and merged["merged_from"] == ["R001", "R002"]
    assert [s["document"] for s in merged["sources"]] == ["a.pdf", "b.pdf"]
    assert merged["flags"] == ["uncapped_liability"]
    assert out[1]["merged_from"] == ["R003"] and len(out[1]["sources"]) == 1


def test_parse_money():
    assert parse_money("2,500,000") == 2_500_000
    assert parse_money("£2.5m") == 2_500_000
    assert parse_money("750k") == 750_000
    assert parse_money(None) is None
    with pytest.raises(CalculationError):
        parse_money("about two million")


# ----- the full Stage 4 run ------------------------------------------------------

GROUND_Z3 = record(category="GRD", title="Ground risk transferred", correlation_group="ground",
                   quote="No compensation event arises from physical conditions within the Site",
                   likelihood=3, impact_min=0.5, impact_likely=2.5, impact_max=8.0, flags=[])
GROUND_SITE = record(category="GRD", title="Made ground found", correlation_group="ground",
                     quote="Made ground up to 3 m deep was found.", likelihood=4, impact_min=0.5,
                     impact_likely=1.5, impact_max=5.0, flags=[])
LIABILITY = record(category="LIA", title="Uncapped liability", correlation_group="liability",
                   quote="The Contractor's total liability to the Employer shall not be limited", likelihood=2,
                   impact_min=1, impact_likely=4, impact_max=20, flags=["uncapped_liability"])
VE = record(type="opportunity", category="VAR", title="Provisional sum may be lower", correlation_group="",
            quote="Provisional sum: diversion of existing water main", likelihood=3, impact_min=0.1,
            impact_likely=0.4, impact_max=1.0, suggested_response="accept", flags=[])


def answers(doc, request):
    if doc is None:   # the duplicate-merging call
        ids = [json.loads(line)["id"] for line in request["messages"][0]["content"].splitlines()
               if line.startswith("{")]
        titles = {json.loads(line)["id"]: json.loads(line)["title"] for line in
                  request["messages"][0]["content"].splitlines() if line.startswith("{")}
        ground = [i for i in ids if titles[i] in ("Ground risk transferred", "Made ground found")]
        opp = [i for i in ids if titles[i] == "Provisional sum may be lower"]
        return response({"groups": [
            {"record_ids": ground, "title": "Unforeseen ground conditions at contractor's risk", "reason": "same exposure"},
            {"record_ids": [ground[0], opp[0]], "title": "bad", "reason": "mixes types"},
        ]})
    data = {"Schedule of Amendments.pdf": [record(), GROUND_Z3, LIABILITY],
            "Site Information.txt": [GROUND_SITE], "Pricing Schedule.xlsx": [VE]}.get(doc, [])
    return response({"records": data, "checklist": []})


def extract_demo(settings, fake):
    plan = plan_extraction(DEMO, settings)
    return run_extraction(plan, settings, client=fake, progress=lambda m: None)


def test_stage4_run_merges_scores_and_simulates(settings):
    fake = FakeClaude(answers)
    ex = extract_demo(settings, fake)
    assert len(ex["records"]) == 5
    out = run_calculation(Path(ex["out_dir"]), settings, 2_500_000, client=fake, progress=lambda m: None)
    a = out["audit"]

    assert a["records_before_merging"] == 5 and a["records_after_merging"] == 4
    assert len(a["merging"]["groups"]) == 1 and a["merging"]["refused"][0]["problem"] == "mixes threats and opportunities"
    merged = next(x for x in a["result"]["records"] if len(x["merged_from"]) == 2)
    assert merged["title"] == "Unforeseen ground conditions at contractor's risk"
    assert {s["document"] for s in merged["sources"]} == {"Schedule of Amendments.pdf", "Site Information.txt"}

    reds = {x["title"] for x in a["result"]["records"] if x["rating"] == "red"}
    assert {"Uncapped delay damages", "Uncapped liability"} <= reds       # flags force red
    c = a["result"]["contingency"]
    assert c["P10"]["pct"] <= c["P50"]["pct"] <= c["P80"]["pct"] <= c["P90"]["pct"]
    assert c["P80"]["money"] == pytest.approx(c["P80"]["pct"] * 25_000)
    assert a["result"]["expected_value"]["opportunities_pct"] < 0

    # merging cost is in the ledger, under the extraction run
    rows = Ledger(Path(settings["budget"]["ledger_file"])).rows()
    merge_rows = [row for row in rows if row["step"] == "Duplicate merging"]
    assert len(merge_rows) == 1 and merge_rows[0]["run_id"] == ex["run_id"]

    wb = openpyxl.load_workbook(out["workbook"])
    assert wb.sheetnames == ["Summary", "Register", "Top contributors", "Cumulative curve", "Merging", "Your edits"]
    reg = wb["Register"]
    assert reg["A2"].value == "ID" and reg["E3"].value == "RED"          # red risks listed first
    assert reg["E3"].fill.fgColor.rgb.endswith(PACK["rating"]["colours"]["red"])
    assert json.loads(Path(out["audit_file"]).read_text())["result"]["simulation"]["iterations"] == 5000


def test_same_records_same_numbers_and_no_merge_option(settings):
    fake = FakeClaude(answers)
    ex = extract_demo(settings, fake)
    first = run_calculation(Path(ex["out_dir"]), settings, 1e6, merge=False, progress=lambda m: None)
    second = run_calculation(Path(ex["out_dir"]), settings, 1e6, merge=False, progress=lambda m: None)
    assert first["out_dir"] != second["out_dir"]                    # never overwritten
    assert first["audit"]["result"] == second["audit"]["result"]    # identical numbers
    assert first["audit"]["records_after_merging"] == 5
    assert first["audit"]["cost_usd"] == 0


def test_edits_in_records_xlsx_are_used(settings):
    fake = FakeClaude(answers)
    ex = extract_demo(settings, fake)
    path = Path(ex["out_dir"]) / "records.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["Risk records"]
    headings = [c.value for c in ws[2]]
    col = {h: i + 1 for i, h in enumerate(headings)}
    ws.cell(row=3, column=col["Likelihood (1-5)"], value=5)
    ws.cell(row=3, column=col["Impact likely %"], value=0.8)   # R001 is the opportunity (max 1%)
    ws.delete_rows(7)                                     # remove the last record
    wb.save(path)

    out = run_calculation(Path(ex["out_dir"]), settings, None, merge=False, progress=lambda m: None)
    edits = out["audit"]["excel_edits"]
    assert {(e["id"], e["field"], e["now"]) for e in edits} >= {("R001", "likelihood", 5), ("R001", "impact_likely", 0.8)}
    assert any(e["field"] == "(whole record)" for e in edits)
    assert out["audit"]["records_before_merging"] == 4
    r1 = next(x for x in out["audit"]["result"]["records"] if x["id"] == "R001")
    assert r1["likelihood"] == 5 and r1["impact_likely"] == 0.8


def test_bad_edit_is_reported_plainly(settings):
    fake = FakeClaude(answers)
    ex = extract_demo(settings, fake)
    path = Path(ex["out_dir"]) / "records.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["Risk records"]
    col = {c.value: i + 1 for i, c in enumerate(ws[2])}
    ws.cell(row=3, column=col["Impact min %"], value=50)  # above most likely
    wb.save(path)
    with pytest.raises(CalculationError, match="row 3 \\(R001\\).*in order"):
        run_calculation(Path(ex["out_dir"]), settings, None, merge=False, progress=lambda m: None)


def test_review_cap_covers_whole_run(settings):
    fake = FakeClaude(answers)
    ex = extract_demo(settings, fake)
    settings["budget"]["per_run_cap_usd"] = ex["summary"]["actual_usd"] + 0.0001   # no room left for merging
    with pytest.raises(CalculationError, match="per-run cap"):
        run_calculation(Path(ex["out_dir"]), settings, None, client=fake, progress=lambda m: None,
                        already_spent_usd=ex["summary"]["actual_usd"])


def test_not_an_extraction_folder(settings, tmp_path):
    with pytest.raises(CalculationError, match="No extraction results"):
        run_calculation(tmp_path, settings, None, merge=False)
