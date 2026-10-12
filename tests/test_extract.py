"""Tests for extraction and quote verification (Stage 3), using a fake Claude."""

import json
from pathlib import Path

import anthropic
import httpx2
import openpyxl
import pytest

from risk_engine.config import PROJECT_ROOT, load_settings
from risk_engine.costs import Ledger
from risk_engine.extract import build_schema, record_problems
from risk_engine.ingest import Document, Section
from risk_engine.packs import load_pack
from risk_engine.pipeline import plan_extraction, run_extraction
from risk_engine.verify import find_quote
from tests.fake_claude import FakeClaude, record, response

DEMO = PROJECT_ROOT / "samples" / "demo_tender"
PACK = load_pack("construction_tender")


@pytest.fixture
def settings(tmp_path):
    s = load_settings()
    s["budget"]["ledger_file"] = str(tmp_path / "ledger.csv")
    s["output_folder"] = str(tmp_path / "outputs")
    return s


# ----- quote verification ----------------------------------------------------

def doc_with(text, locations=None):
    section = Section("d.pdf", "Z3 Ground conditions (pages 1-2)", text,
                      locations=locations or [(0, "page 1")])
    return Document("d.pdf", Path("d.pdf"), "PDF", [section])


@pytest.mark.parametrize("quote,method", [
    ("Contractor's total liability shall not be limited", "exact"),        # curly apostrophe, line break
    ("the obligation is fit-for-purpose", "spacing-tolerant"),             # broken word + kept hyphen
    ("is fit-for-purpose", "exact"),
    ("The Agreement was previously signed", "spacing-tolerant"),           # PDF spacing faults
    ("Shall not be limited and the", "exact"),                              # first letter capitalised
])
def test_quotes_that_verify(quote, method):
    d = doc_with("The Contractor’s total liability\nshall not be limited and the obliga-\ntion is fit-for-\n"
                 "purpose. The A greement was previo usly signed.")
    m = find_quote(quote, d)
    assert m.verified and m.method == method


@pytest.mark.parametrize("quote,reason", [
    ("The Contractor's total liability shall be limited", "not found"),   # one word changed
    ("The Contractor's total liability ... not limited", "'...'"),
    ("not limited", "too short"),
    ("The Contractor's Total liability shall not be limited", "not found"), # case changed mid-quote
    ("", "no quote"),
])
def test_quotes_that_fail(quote, reason):
    d = doc_with("The Contractor's total liability shall not be limited.")
    m = find_quote(quote, d)
    assert not m.verified and reason in m.reason


def test_verified_quote_gets_exact_page_from_code():
    d = doc_with("First page text here.\nSecond page: the price is fixed for the duration.",
                 locations=[(0, "page 1"), (22, "page 2")])
    m = find_quote("the price is fixed for the duration", d)
    assert m.ref == "Z3 Ground conditions (page 2)"


def test_quote_running_over_two_sections():
    a = Section("d", "Z1 Order (page 1)", "These Z clauses take", locations=[(0, "page 1")])
    b = Section("d", "Z2 Liability (page 1)", "precedence over all other documents.", locations=[(0, "page 1")])
    m = find_quote("These Z clauses take precedence over all", Document("d", Path("d"), "PDF", [a, b]))
    assert m.verified and m.section is a


# ----- what Claude is asked for ------------------------------------------------

def test_schema_comes_from_the_pack():
    schema = build_schema(PACK)
    rec = schema["properties"]["records"]["items"]
    assert rec["properties"]["category"]["enum"] == PACK.taxonomy_codes
    assert rec["properties"]["suggested_response"]["enum"] == PACK.config["responses"]
    assert set(rec["required"]) == set(rec["properties"])
    ids = schema["properties"]["checklist"]["items"]["properties"]["id"]["enum"]
    assert ids == [c["id"] for c in PACK.checklist]

    def objects(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                yield node
            for v in node.values():
                yield from objects(v)
        elif isinstance(node, list):
            for v in node:
                yield from objects(v)
    assert all(o.get("additionalProperties") is False for o in objects(schema))


def test_record_checks():
    assert record_problems(record(), PACK) == []
    assert "impact range is not in order" in record_problems(record(impact_min=5, impact_likely=1), PACK)[0]
    assert "category" in record_problems(record(category="XYZ"), PACK)[0]
    assert "likelihood" in record_problems(record(likelihood=9), PACK)[0]


# ----- a full run with the fake Claude ------------------------------------------

def demo_answers(doc, request):
    if doc == "Schedule of Amendments.pdf":
        return response({
            "records": [
                record(),
                record(title="Invented clause", quote="The Employer may withhold all payments indefinitely."),
                record(title="Bad range", category="LIA", quote="liability for indirect or consequential loss is excluded",
                       impact_min=4, impact_likely=1, impact_max=2, flags=[]),
            ],
            "checklist": [{"id": "ld_rate_cap", "answer": "GBP 15,000 per week, no cap",
                           "quote": "Delay damages are GBP 15,000 per week.", "source_ref": "Z6"}],
        }, cache_read=3000)
    return response(cache_write=3000 if doc.startswith("Instructions") else 0,
                    cache_read=0 if doc.startswith("Instructions") else 3000)


def test_full_run_on_demo_pack(settings):
    plan = plan_extraction(DEMO, settings)
    fake = FakeClaude(demo_answers)
    r = run_extraction(plan, settings, client=fake, progress=lambda m: None)

    # every piece was sent once; the instructions are marked for caching
    assert len(fake.calls) == len(plan.chunks) == 4
    first = fake.calls[0]
    assert first["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert first["output_config"]["format"]["type"] == "json_schema"
    assert first["model"] == settings["models"]["extraction"]
    assert all(c["system"] == first["system"] for c in fake.calls)   # identical, so it caches

    # one genuine record accepted, with the reference set by code
    assert [x["id"] for x in r["records"]] == ["R001"]
    rec = r["records"][0]
    assert rec["source_ref"] == "Z6 Delay damages (page 2)" and rec["quote_verified"]
    assert rec["source_doc"] == "Schedule of Amendments.pdf"

    # the invented quote and the impossible range are excluded, with reasons
    reasons = {x["title"]: x["problems"] for x in r["excluded"]}
    assert "quote not verified" in reasons["Invented clause"][0]
    assert any("not in order" in p for p in reasons["Bad range"])

    # checklist: answered where a verified quote exists, otherwise reported as not found
    status = {c["id"]: c["status"] for c in r["checklist"]}
    assert status["ld_rate_cap"] == "answered"
    assert status["liability_cap"] == "not found in the documents"
    assert len(status) == len(PACK.checklist)

    # spending: one ledger line per call with the real cost, matching the summary
    rows = [row for row in Ledger(Path(settings["budget"]["ledger_file"])).rows() if row["mode"] == "live"]
    assert len(rows) == 4
    assert sum(float(row["actual_usd"]) for row in rows) == pytest.approx(r["summary"]["actual_usd"], abs=1e-5)
    assert r["summary"]["actual_usd"] > 0

    # outputs written
    audit = json.loads(Path(r["audit_file"]).read_text())
    assert audit["models"]["extraction"] == settings["models"]["extraction"]
    assert "Role" in audit["prompts"]["system"]
    wb = openpyxl.load_workbook(r["records_file"])
    assert wb.sheetnames == ["Run summary", "Risk records", "Excluded", "Checklist", "Pieces reviewed"]
    assert wb["Risk records"]["A3"].value == "R001"


def test_cap_stops_run_before_spending(settings):
    plan = plan_extraction(DEMO, settings)
    plan.cap_usd = 0.0001
    fake = FakeClaude(demo_answers)
    r = run_extraction(plan, settings, client=fake, progress=lambda m: None)
    assert fake.calls == []
    assert "per-run cap" in r["summary"]["stopped_because"]
    assert r["summary"]["actual_usd"] == 0
    assert Path(r["records_file"]).exists()


def test_rejected_key_stops_with_plain_message(settings):
    def answers(doc, request):
        req = httpx2.Request("POST", "https://api.anthropic.com/v1/messages")
        return anthropic.AuthenticationError("invalid x-api-key", response=httpx2.Response(401, request=req), body=None)
    plan = plan_extraction(DEMO, settings)
    fake = FakeClaude(answers)
    r = run_extraction(plan, settings, client=fake, progress=lambda m: None)
    assert len(fake.calls) == 1                     # stopped after the first failure
    assert r["summary"]["stopped_because"].startswith("The API key was rejected")
    assert r["summary"]["pieces_reviewed"] == 0


def test_refusal_and_cut_off_answers_are_reported(settings):
    def answers(doc, request):
        if doc.startswith("Pricing"):
            return response(stop_reason="refusal", text="")
        if doc.startswith("Site"):
            return response(stop_reason="max_tokens", text='{"records": [')
        return response()
    plan = plan_extraction(DEMO, settings)
    r = run_extraction(plan, settings, client=FakeClaude(answers), progress=lambda m: None)
    status = {p.doc: p.status for p in r["pieces"]}
    assert status["Pricing Schedule.xlsx"] == "refused"
    assert status["Site Information.txt"] == "too_long"
    assert r["summary"]["pieces_not_reviewed"] == 2
    assert not r["summary"]["stopped_because"]


def test_real_sdk_request_and_streamed_reply(settings):
    """Drive the real anthropic SDK against a fake network connection that
    answers like the Claude API, to check the request on the wire and that the
    streamed reply is read correctly."""
    answer = json.dumps({"records": [record()], "checklist": []})
    events = [
        ("message_start", {"type": "message_start", "message": {
            "id": "msg_1", "type": "message", "role": "assistant", "model": "claude-sonnet-5-5", "content": [],
            "stop_reason": None, "stop_sequence": None,
            "usage": {"input_tokens": 900, "output_tokens": 1, "cache_creation_input_tokens": 3100,
                      "cache_read_input_tokens": 0}}}),
        ("content_block_start", {"type": "content_block_start", "index": 0,
                                 "content_block": {"type": "text", "text": ""}}),
        ("content_block_delta", {"type": "content_block_delta", "index": 0,
                                 "delta": {"type": "text_delta", "text": answer}}),
        ("content_block_stop", {"type": "content_block_stop", "index": 0}),
        ("message_delta", {"type": "message_delta", "delta": {"stop_reason": "end_turn", "stop_sequence": None},
                           "usage": {"output_tokens": 1800}}),
        ("message_stop", {"type": "message_stop"}),
    ]
    sse = "".join(f"event: {name}\ndata: {json.dumps(data)}\n\n" for name, data in events)
    sent = []

    def handler(request):
        sent.append(json.loads(request.content))
        return httpx2.Response(200, headers={"content-type": "text/event-stream", "request-id": "req_wire"},
                               content=sse.encode())

    client = anthropic.Anthropic(api_key="test-key", http_client=httpx2.Client(transport=httpx2.MockTransport(handler)))
    plan = plan_extraction(DEMO, settings)
    plan.chunks = [c for c in plan.chunks if c.doc == "Schedule of Amendments.pdf"]
    r = run_extraction(plan, settings, client=client, progress=lambda m: None)

    body = sent[0]
    assert body["model"] == "claude-sonnet-5-5" and body["stream"] is True
    assert body["system"][0]["cache_control"] == {"type": "ephemeral"}
    assert body["output_config"]["effort"] == "medium"
    assert body["output_config"]["format"]["type"] == "json_schema"
    assert "thinking" not in body and "temperature" not in body
    piece = r["pieces"][0]
    assert piece.status == "ok" and piece.request_id == "req_wire"
    assert piece.usage == {"input": 900, "cache_write": 3100, "cache_read": 0, "output": 1800}
    assert [x["id"] for x in r["records"]] == ["R001"]
