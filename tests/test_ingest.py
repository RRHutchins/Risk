"""Tests for document reading and splitting (Stage 2)."""

import docx
import openpyxl
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from risk_engine.config import PROJECT_ROOT
from risk_engine.ingest import HeadingDetector, chunk_documents, read_docx, read_folder
from risk_engine.packs import load_pack

DEMO = PROJECT_ROOT / "samples" / "demo_tender"
PACK_PATTERNS = load_pack("construction_tender").heading_patterns


@pytest.fixture(scope="module")
def demo():
    docs, skipped = read_folder(DEMO, extra_heading_patterns=PACK_PATTERNS)
    return {d.name: d for d in docs}, skipped


def find(doc, phrase):
    """The section containing a phrase, and the page/row/line of that phrase."""
    for s in doc.sections:
        i = s.text.find(phrase)
        if i >= 0:
            return s, s.location_at(i)
    raise AssertionError(f"'{phrase}' not found in {doc.name}")


# ----- heading detection ---------------------------------------------------

@pytest.mark.parametrize("line,label", [
    ("Z6 Delay damages", "Z6"), ("4.2 The Contractor shall", "4.2"), ("2. Clarifications", "2"),
    ("Clause 7.1 Payment", "Clause 7.1"), ("SCHEDULE 2 - Pricing", "SCHEDULE 2"),
    ("Amendment 14: Liability", "Amendment 14"), ("Appendix A", "Appendix A"),
    ("10 December 2026", None), ("64 Nr piles", None), ("5% of the Price", None),
    ("2026 Edition", None), ("Zone A", None), ("The Contractor shall", None),
])
def test_heading_detection(line, label):
    assert HeadingDetector(PACK_PATTERNS).match(line) == label


def test_engine_alone_does_not_know_nec_clauses():
    assert HeadingDetector().match("Z6 Delay damages") is None


# ----- the demo pack -------------------------------------------------------

def test_all_demo_files_read(demo):
    docs, skipped = demo
    assert len(docs) == 4 and skipped == []


def test_pdf_split_by_clause_with_pages(demo):
    pdf = demo[0]["Schedule of Amendments.pdf"]
    assert [s.clause for s in pdf.sections] == [None] + [f"Z{i}" for i in range(1, 10)]
    refs = {s.clause: s.ref for s in pdf.sections}
    assert refs["Z2"] == "Z2 Liability (page 1)"
    assert refs["Z3"] == "Z3 Ground conditions (pages 1-2)"
    assert refs["Z6"] == "Z6 Delay damages (page 2)"
    assert refs["Z9"] == "Z9 Security (page 3)"


def test_pdf_quote_maps_to_exact_page(demo):
    pdf = demo[0]["Schedule of Amendments.pdf"]
    section, page = find(pdf, "ground conditions, including any obstructions")
    assert section.clause == "Z3" and page == "page 1"
    section, page = find(pdf, "No compensation event arises")
    assert section.clause == "Z3" and page == "page 2"
    _, page = find(pdf, "The total of delay damages is not capped.")
    assert page == "page 2"


def test_pdf_headers_and_footers_removed(demo):
    pdf = demo[0]["Schedule of Amendments.pdf"]
    text = "\n".join(s.text for s in pdf.sections)
    assert "Page 2 of 3" not in text
    assert "Tender Issue" not in text
    assert any("header/footer" in n for n in pdf.notes)


def test_word_split_with_pages_and_parent_headings(demo):
    word = demo[0]["Instructions to Tenderers.docx"]
    refs = [s.ref for s in word.sections]
    assert refs[1] == "1. Tender return (page 1)"
    assert "clause 1.2 in '1. Tender return' (page 1)" in refs
    assert "3. Form of contract (page 2)" in refs
    assert "clause 4.2 in '4. Evaluation' (page 2)" in refs


def test_word_table_kept_in_place(demo):
    word = demo[0]["Instructions to Tenderers.docx"]
    section, page = find(word, "Completion date | 30 November 2027")
    assert section.heading == "5. Key dates" and page == "page 2"


def test_excel_blocks_and_rows(demo):
    xl = demo[0]["Pricing Schedule.xlsx"]
    assert [s.ref for s in xl.sections] == [
        "sheet 'Pricing' (rows 1-50)", "sheet 'Pricing' (rows 51-56)", "sheet 'Notes' (rows 1-2)"]
    _, row = find(xl, "Provisional sum: diversion of existing water main")
    assert row == "row 55"


def test_text_file_sections_and_lines(demo):
    txt = demo[0]["Site Information.txt"]
    assert [s.ref for s in txt.sections][1:] == [
        "1 Site description (lines 3-5)", "2 Existing buildings (lines 7-8)", "3 Access (lines 10-12)"]
    _, line = find(txt, "No asbestos survey")
    assert line == "line 8"


def test_no_text_is_lost_in_chunks(demo):
    docs = list(demo[0].values())
    chunks = chunk_documents(docs, max_chars=1500)
    joined = {d.name: "\n".join(c.text for c in chunks if c.doc == d.name) for d in docs}
    for d in docs:
        for s in d.sections:
            for line in s.text.split("\n"):
                assert line in joined[d.name]


# ----- awkward files -------------------------------------------------------

def test_word_without_saved_layout_has_no_page_numbers(tmp_path):
    d = docx.Document()
    d.add_heading("1. Scope", 1)
    d.add_paragraph("1.1 The works comprise a new roof.")
    d.save(tmp_path / "plain.docx")
    doc = read_docx(tmp_path / "plain.docx", HeadingDetector())
    assert [s.ref for s in doc.sections] == ["1. Scope", "clause 1.1 in '1. Scope'"]
    assert any("page numbers" in n for n in doc.notes)


def test_scanned_pdf_is_warned(tmp_path):
    c = canvas.Canvas(str(tmp_path / "scan.pdf"), pagesize=A4)
    c.rect(50, 50, 200, 200, fill=1)  # a picture, no text
    c.showPage()
    c.save()
    docs, _ = read_folder(tmp_path)
    assert docs[0].sections == []
    assert "scanned" in docs[0].warnings[0]


def test_partly_scanned_pdf_names_the_pages(tmp_path):
    c = canvas.Canvas(str(tmp_path / "mixed.pdf"), pagesize=A4)
    c.drawString(50, 700, "1 Introduction")
    c.showPage()
    c.rect(50, 50, 200, 200, fill=1)
    c.showPage()
    c.save()
    docs, _ = read_folder(tmp_path)
    assert "pages 2" in docs[0].warnings[0]


def test_damaged_files_get_plain_messages(tmp_path):
    (tmp_path / "broken.docx").write_bytes(b"this is not a word file")
    (tmp_path / "broken.pdf").write_bytes(b"this is not a pdf")
    (tmp_path / "broken.xlsx").write_bytes(b"this is not a spreadsheet")
    docs, skipped = read_folder(tmp_path)
    assert docs == []
    reasons = {s.name: s for s in skipped}
    assert "not a real Word file" in reasons["broken.docx"].reason
    assert "not a real PDF" in reasons["broken.pdf"].reason
    assert "not a real Excel file" in reasons["broken.xlsx"].reason
    for s in skipped:
        assert "Error" not in s.reason          # no technical names shown to the user
        assert s.detail                          # but kept for the audit file


def test_old_formats_get_advice(tmp_path):
    (tmp_path / "old.doc").write_bytes(b"x")
    (tmp_path / "old.xls").write_bytes(b"x")
    (tmp_path / "site.jpg").write_bytes(b"x")
    (tmp_path / "~$lockfile.docx").write_bytes(b"x")
    _, skipped = read_folder(tmp_path)
    reasons = {s.name: s.reason for s in skipped}
    assert "Save As" in reasons["old.doc"] and "Save As" in reasons["old.xls"]
    assert "pictures" in reasons["site.jpg"]
    assert "~$lockfile.docx" not in reasons


def test_subfolders_are_read_and_named(tmp_path):
    (tmp_path / "Contract").mkdir()
    (tmp_path / "Contract" / "notes.txt").write_text("1 Scope\nNew roof.")
    docs, _ = read_folder(tmp_path)
    assert docs[0].name.replace("\\", "/") == "Contract/notes.txt"
    assert docs[0].sections[0].doc == docs[0].name


def test_hidden_excel_sheet_is_noted(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["visible"])
    hidden = wb.create_sheet("Secret")
    hidden.append(["hidden rates"])
    hidden.sheet_state = "hidden"
    wb.save(tmp_path / "book.xlsx")
    docs, _ = read_folder(tmp_path)
    assert any("Secret" in n for n in docs[0].notes)
    assert any("hidden rates" in s.text for s in docs[0].sections)
