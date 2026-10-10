"""Tests for settings, domain packs, document reading and the dry run."""

import json
from pathlib import Path

import openpyxl
import pytest

from risk_engine.config import PROJECT_ROOT, ConfigError, load_settings
from risk_engine.costs import Ledger
from risk_engine.dryrun import dry_run, ingest_only
from risk_engine.ingest import Document, Section, chunk_documents
from risk_engine.packs import load_pack

DEMO = PROJECT_ROOT / "samples" / "demo_tender"


def test_settings_load():
    s = load_settings()
    assert s["budget"]["per_run_cap_usd"] == 3.00
    assert s["pricing"]["claude-sonnet-5-5"]["cache_read"] == 0.10
    assert s["simulation"]["iterations"] == 10000
    for model in s["models"].values():
        assert model in s["pricing"], f"no price for {model}"


def test_construction_pack_loads():
    pack = load_pack("construction_tender")
    assert len(pack.taxonomy_codes) == 16
    assert len(pack.checklist) == 10
    assert 5 <= len(pack.examples) <= 10
    for ex in pack.examples:
        assert ex["category"] in pack.taxonomy_codes


def test_missing_pack_is_reported():
    with pytest.raises(ConfigError):
        load_pack("no_such_pack")


def test_chunks_respect_size_and_keep_documents_apart():
    long_text = "\n".join(f"Line {i} " + "x" * 80 for i in range(100))
    docs = [Document("a.txt", Path("a.txt"), "Text", [Section("a.txt", "whole document", long_text)]),
            Document("b.txt", Path("b.txt"), "Text", [Section("b.txt", "whole document", "short")])]
    chunks = chunk_documents(docs, max_chars=2000)
    assert all(len(c.text) <= 2000 for c in chunks)
    assert chunks[-1].doc == "b.txt" and all(c.doc == "a.txt" for c in chunks[:-1])
    # no text lost
    assert sum(c.text.count("Line ") for c in chunks if c.doc == "a.txt") == 100


def test_dry_run_end_to_end(tmp_path):
    settings = load_settings()
    settings["budget"]["ledger_file"] = str(tmp_path / "ledger.csv")
    settings["output_folder"] = str(tmp_path / "outputs")
    r = dry_run(DEMO, settings)
    assert len(r["documents"]) == 4
    assert r["estimate"]["total_usd"] > 0
    assert r["within_cap"]
    audit = json.loads(Path(r["audit_file"]).read_text())
    assert audit["run_id"] == r["run_id"]
    assert "extraction_context" in audit["prompts"]
    assert (tmp_path / "ledger.csv").exists()


def test_runs_started_together_get_separate_ids_and_files(tmp_path):
    settings = load_settings()
    settings["budget"]["ledger_file"] = str(tmp_path / "ledger.csv")
    settings["output_folder"] = str(tmp_path / "outputs")
    runs = [dry_run(DEMO, settings) for _ in range(5)]
    ids = {r["run_id"] for r in runs}
    assert len(ids) == 5
    assert all(Path(r["audit_file"]).exists() for r in runs)
    assert Ledger(tmp_path / "ledger.csv").run_count() == 5


def test_ingest_writes_sections_workbook(tmp_path):
    settings = load_settings()
    settings["output_folder"] = str(tmp_path / "outputs")
    r = ingest_only(DEMO, settings)
    wb = openpyxl.load_workbook(r["sections_file"])
    rows = list(wb["Sections"].iter_rows(values_only=True))
    assert rows[0][3] == "Reference"
    assert any(row[3] == "Z3 Ground conditions (pages 1-2)" for row in rows)


def test_dry_run_is_repeatable(tmp_path):
    settings = load_settings()
    settings["budget"]["ledger_file"] = str(tmp_path / "ledger.csv")
    settings["output_folder"] = str(tmp_path / "outputs")
    a, b = dry_run(DEMO, settings), dry_run(DEMO, settings)
    assert a["estimate"] == b["estimate"]
    assert a["tokens"] == b["tokens"]
