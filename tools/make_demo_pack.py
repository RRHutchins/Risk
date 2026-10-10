"""Builds the small fictional demo tender pack in samples/demo_tender/.

Developer tool only: the owner does not need to run this. The files it makes
are already included in the project. Needs reportlab (pip install reportlab).
All wording is invented; none is copied from any standard form of contract.

The pack is designed to test document reading (Stage 2):
  - the PDF has a header and footer on every page, and clause Z3 runs
    across a page break
  - the Word file has headings, numbered clauses, a table, and page layout
  - the spreadsheet is long enough to be split into blocks of rows
"""

from pathlib import Path

import docx
import openpyxl
from docx.oxml import OxmlElement
from docx.enum.text import WD_BREAK
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "samples" / "demo_tender"
PROJECT = "Riverside Community Centre (fictional demonstration project)"
EMPLOYER = "Northvale Borough Council (fictional)"


# ---------------------------------------------------------------------------
# Word: Instructions to Tenderers
# ---------------------------------------------------------------------------

def _page_break_before(paragraph):
    """Start this paragraph on a new page, recording the break the way Word
    does when it saves a file (a hard break plus a 'last rendered page break'
    marker), so page numbers can be read back."""
    run = paragraph.insert_paragraph_before().add_run()
    run.add_break(WD_BREAK.PAGE)
    marker_run = paragraph.runs[0]._r if paragraph.runs else paragraph.add_run()._r
    marker_run.insert(0, OxmlElement("w:lastRenderedPageBreak"))


def instructions_to_tenderers():
    d = docx.Document()
    d.add_heading("Instructions to Tenderers", 0)
    d.add_paragraph(f"Project: {PROJECT}")
    d.add_paragraph(f"Employer: {EMPLOYER}")

    d.add_heading("1. Tender return", 1)
    d.add_paragraph("1.1 Tenders must be returned through the Employer's e-tendering portal by 12:00 noon "
                    "on 31 October 2026.")
    d.add_paragraph("1.2 Late tenders will not be considered.")

    d.add_heading("2. Clarifications", 1)
    d.add_paragraph("2.1 Clarification questions must be submitted no later than 10 working days before the "
                    "tender return date. Questions received after this date will not be answered.")
    d.add_paragraph("2.2 All questions and answers will be shared with every tenderer without naming the "
                    "tenderer who asked.")

    h3 = d.add_heading("3. Form of contract", 1)
    _page_break_before(h3)
    d.add_paragraph("3.1 The contract will be the NEC4 Engineering and Construction Contract, Option A, "
                    "amended by the Schedule of Amendments included in this tender pack.")
    d.add_paragraph("3.2 The Contractor's design is limited to the mechanical and electrical services.")

    d.add_heading("4. Evaluation", 1)
    d.add_paragraph("4.1 Tenders will be evaluated 70% on price and 30% on quality.")
    d.add_paragraph("4.2 Qualified tenders may be rejected at the Employer's sole discretion.")

    d.add_heading("5. Key dates", 1)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Event", "Date"
    for event, date in [("Tender return", "31 October 2026"), ("Access date", "5 January 2027"),
                        ("Completion date", "30 November 2027")]:
        row = t.add_row()
        row.cells[0].text, row.cells[1].text = event, date
    d.add_paragraph("5.1 The Employer does not intend to agree any extension to the Completion Date at tender stage.")
    d.save(OUT / "Instructions to Tenderers.docx")


# ---------------------------------------------------------------------------
# PDF: Schedule of Amendments (drawn line by line so page breaks are exact)
# ---------------------------------------------------------------------------

AMENDMENT_PAGES = [
    [
        ("title", "Schedule of Amendments (Option Z clauses)"),
        ("text", PROJECT),
        ("gap", ""),
        ("head", "Z1 Order of precedence"),
        ("text", "If there is any conflict between the documents, these Z clauses take precedence over"),
        ("text", "all other documents, and the Scope takes precedence over the Contractor's tender."),
        ("gap", ""),
        ("head", "Z2 Liability"),
        ("text", "The Contractor's total liability to the Employer shall not be limited, save that"),
        ("text", "liability for indirect or consequential loss is excluded."),
        ("gap", ""),
        ("head", "Z3 Ground conditions"),
        ("text", "The Contractor is deemed to have inspected the Site and satisfied itself as to the"),
        ("text", "ground conditions, including any obstructions, contamination and groundwater."),
    ],
    [
        ("text", "No compensation event arises from physical conditions within the Site, whether or"),
        ("text", "not an experienced contractor would have allowed for them."),
        ("gap", ""),
        ("head", "Z4 Payment"),
        ("text", "The Employer pays within 45 days of the assessment date. Retention is 5% of the"),
        ("text", "Price for Work Done to Date, half of which is released at Completion."),
        ("gap", ""),
        ("head", "Z5 Notification"),
        ("text", "A compensation event not notified by the Contractor within 4 weeks of becoming"),
        ("text", "aware of it is time-barred."),
        ("gap", ""),
        ("head", "Z6 Delay damages"),
        ("text", "Delay damages are GBP 15,000 per week. The total of delay damages is not capped."),
    ],
    [
        ("head", "Z7 Termination"),
        ("text", "The Employer may terminate the Contractor's obligation to Provide the Works at any"),
        ("text", "time for its convenience. The Contractor is not entitled to loss of profit on work"),
        ("text", "not carried out."),
        ("gap", ""),
        ("head", "Z8 Price"),
        ("text", "The Prices are fixed for the duration of the contract and are not subject to"),
        ("text", "adjustment for inflation."),
        ("gap", ""),
        ("head", "Z9 Security"),
        ("text", "The Contractor provides a performance bond of 10% of the total of the Prices and a"),
        ("text", "parent company guarantee."),
    ],
]


def schedule_of_amendments():
    c = canvas.Canvas(str(OUT / "Schedule of Amendments.pdf"), pagesize=A4)
    width, height = A4
    total = len(AMENDMENT_PAGES)
    for number, lines in enumerate(AMENDMENT_PAGES, start=1):
        c.setFont("Helvetica", 8)
        c.drawString(50, height - 35, "Riverside Community Centre - Schedule of Amendments - Tender Issue")
        c.drawRightString(width - 50, 30, f"Page {number} of {total}")
        y = height - 80
        for kind, text in lines:
            if kind == "title":
                c.setFont("Helvetica-Bold", 15)
            elif kind == "head":
                c.setFont("Helvetica-Bold", 11)
            else:
                c.setFont("Helvetica", 10)
            if text:
                c.drawString(50, y, text)
            y -= 22 if kind == "title" else 15
        c.showPage()
    c.save()


# ---------------------------------------------------------------------------
# Excel: Pricing Schedule (long enough to be split into blocks)
# ---------------------------------------------------------------------------

def pricing_schedule():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws.append(["Item", "Description", "Unit", "Quantity", "Rate (GBP)", "Total (GBP)"])
    ws.append(["1.1", "Site set-up and welfare", "item", 1, None, None])
    ws.append(["2.1", "Demolition of existing pavilion (asbestos survey not yet carried out)", "item", 1, None, None])
    for i in range(1, 51):
        ws.append([f"3.{i}", f"Pile type P{i % 4 + 1}, 450 mm diameter, bored to 12 m", "nr", 2, None, None])
    ws.append(["4.1", "Steel frame", "t", 85, None, None])
    ws.append(["5.1", "Provisional sum: diversion of existing water main", "item", 1, 40000, 40000])
    ws.append(["6.1", "Contractor's design: mechanical and electrical services (fit for purpose)", "item", 1, None, None])
    notes = wb.create_sheet("Notes")
    notes.append(["Tenderers must price all items. Unpriced items are deemed included elsewhere."])
    notes.append(["The site investigation report (2019) is provided for information only."])
    wb.save(OUT / "Pricing Schedule.xlsx")


# ---------------------------------------------------------------------------
# Text: Site Information
# ---------------------------------------------------------------------------

def site_notes():
    (OUT / "Site Information.txt").write_text(
        "Site Information\n"
        "\n"
        "1 Site description\n"
        "The site is a former industrial yard. A site investigation was carried out in 2019 with three\n"
        "boreholes. Made ground up to 3 m deep was found.\n"
        "\n"
        "2 Existing buildings\n"
        "Asbestos may be present in the existing pavilion. No asbestos survey has been carried out.\n"
        "\n"
        "3 Access\n"
        "Access is from Mill Lane only, which is shared with a primary school. No deliveries between\n"
        "08:15-09:00 and 15:00-15:45 on school days.\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    instructions_to_tenderers()
    schedule_of_amendments()
    pricing_schedule()
    site_notes()
    print(f"Demo pack written to {OUT}")
